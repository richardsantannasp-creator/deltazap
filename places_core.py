#!/usr/bin/env python3
"""
Núcleo do Extrator de Contatos — Google Places API (New).

- Cliente da API com novas tentativas automáticas (429/5xx) e contagem de chamadas
- Busca rápida (até 60 por nicho) ou busca em GRADE: divide a área da cidade em
  quadrantes e subdivide os que voltam cheios (60), passando do limite de 60
- Normalização de telefones com a biblioteca phonenumbers (se instalada)
- Exportação para Excel/CSV

Uso pela linha de comando (sem a interface):
    export GOOGLE_PLACES_API_KEY="sua_chave"
    python extrator_google_places.py "dentista" "Guarulhos SP"
    python extrator_google_places.py "dentista;pet shop" "Guarulhos SP" --grade --max-chamadas 120 --sites
"""

import argparse
import csv
import hashlib
import random
import re
import sys
import time

import requests

try:
    import phonenumbers
    from phonenumbers import PhoneNumberType
except ImportError:  # funciona sem, com uma regra mais simples
    phonenumbers = None

API_URL = "https://places.googleapis.com/v1/places:searchText"
FIELD_MASK = ",".join([
    "places.id", "places.displayName", "places.formattedAddress",
    "places.nationalPhoneNumber", "places.internationalPhoneNumber",
    "places.websiteUri", "places.googleMapsUri", "places.rating",
    "places.userRatingCount", "places.businessStatus", "places.location",
    "nextPageToken",
])
FIELD_MASK_CIDADE = "places.id,places.displayName,places.formattedAddress,places.viewport"
LIMITE_POR_BUSCA = 60  # teto do Google por busca de texto (3 páginas de 20)

ORIGEM_SITE = "Site da empresa"
ORIGEM_CELULAR = "Celular do Google"

# (rótulo na planilha, chave interna)
COLUNAS = [
    ("Nome", "nome"), ("WhatsApp", "whatsapp"), ("Origem do WhatsApp", "whatsapp_origem"),
    ("Link WhatsApp", "link_whatsapp"), ("Telefone (Google)", "telefone"), ("Tipo do telefone", "tipo"),
    ("E-mail", "email"), ("Instagram", "instagram"), ("Endereço", "endereco"), ("Site", "site"),
    ("Avaliação", "avaliacao"), ("Nº avaliações", "n_avaliacoes"), ("Google Maps", "maps"),
    ("Nicho", "termo"), ("Status", "status"), ("Notas", "notas"),
]


class ErroAPI(Exception):
    pass


# ---------------------------------------------------------------- telefones
def _regra_simples(numero):
    digitos = re.sub(r"\D", "", numero).lstrip("0")
    if not digitos.startswith("55"):
        digitos = "55" + digitos
    resto = digitos[2:]
    if len(resto) == 11 and resto[2] == "9":
        return "+" + digitos, "Celular"
    if len(resto) == 10 and resto[2] in "2345":
        return "+" + digitos, "Fixo"
    return "+" + digitos, "Desconhecido"


def normalizar_telefone_br(numero):
    """Retorna (E.164, tipo). tipo: 'Celular', 'Fixo', 'Desconhecido' ou '' se vazio."""
    if not numero or not re.search(r"\d", str(numero)):
        return "", ""
    numero = str(numero).strip()
    if numero.startswith("+") and not re.sub(r"[^\d+]", "", numero).startswith("+55"):
        return "+" + re.sub(r"\D", "", numero), "Desconhecido"  # número de outro país
    if not numero.startswith("+"):
        numero = re.sub(r"\D", "", numero).lstrip("0")  # "011 9..." -> "119..."
        if len(numero) in (12, 13) and numero.startswith("55"):
            numero = "+" + numero
    if phonenumbers:
        try:
            n = phonenumbers.parse(numero, "BR")
            if phonenumbers.is_valid_number(n):
                e164 = phonenumbers.format_number(n, phonenumbers.PhoneNumberFormat.E164)
                t = phonenumbers.number_type(n)
                if t == PhoneNumberType.MOBILE:
                    return e164, "Celular"
                if t == PhoneNumberType.FIXED_LINE:
                    return e164, "Fixo"
                if t == PhoneNumberType.FIXED_LINE_OR_MOBILE:
                    return _regra_simples(e164)
                return e164, "Desconhecido"
        except phonenumbers.NumberParseException:
            pass
    return _regra_simples(numero)


def link_whatsapp(e164):
    return f"https://wa.me/{e164.lstrip('+')}" if e164 else ""


# ---------------------------------------------------------------- cliente
class PlacesClient:
    def __init__(self, api_key, demo=False):
        self.api_key = api_key
        self.demo = demo
        self.chamadas = 0
        self._demo = DemoPlaces() if demo else None

    def _post(self, corpo, mascara):
        self.chamadas += 1
        if self.demo:
            time.sleep(0.05)
            return self._demo.search(corpo, mascara)
        headers = {"Content-Type": "application/json", "X-Goog-Api-Key": self.api_key,
                   "X-Goog-FieldMask": mascara}
        for tentativa in range(4):
            try:
                r = requests.post(API_URL, headers=headers, json=corpo, timeout=30)
            except requests.RequestException as e:
                if tentativa == 3:
                    raise ErroAPI(f"Falha de conexão com o Google: {e}")
                time.sleep(2 ** tentativa)
                continue
            if r.status_code == 200:
                return r.json()
            if r.status_code in (429, 500, 502, 503, 504) and tentativa < 3:
                time.sleep(2 ** tentativa + random.random())
                continue
            try:
                msg = r.json().get("error", {}).get("message", r.text)
            except ValueError:
                msg = r.text
            raise ErroAPI(f"Erro {r.status_code} da Places API: {msg[:300]}")

    def area_da_cidade(self, cidade):
        """Retângulo (viewport) da cidade, usado como ponto de partida da grade."""
        d = self._post({"textQuery": cidade, "languageCode": "pt-BR", "regionCode": "BR", "pageSize": 1},
                       FIELD_MASK_CIDADE)
        ps = d.get("places") or []
        if not ps or "viewport" not in ps[0]:
            raise ErroAPI(f'Não encontrei a área de "{cidade}". Tente no formato "Cidade UF".')
        return ps[0]["viewport"]

    def buscar_texto(self, consulta, retangulo=None, maximo=LIMITE_POR_BUSCA, cancelado=None):
        """Retorna (lugares, completo). completo=False se foi interrompido no meio."""
        corpo = {"textQuery": consulta, "languageCode": "pt-BR", "regionCode": "BR", "pageSize": 20}
        if retangulo:
            corpo["locationRestriction"] = {"rectangle": retangulo}
        lugares = []
        while len(lugares) < maximo:
            d = self._post(corpo, FIELD_MASK)
            lugares.extend(d.get("places", []))
            token = d.get("nextPageToken")
            if not token:
                break
            if cancelado and cancelado():
                return lugares[:maximo], False
            corpo["pageToken"] = token
            if not self.demo:
                time.sleep(1.5)
        return lugares[:maximo], True


# ---------------------------------------------------------------- grade
def dividir(r):
    lo, hi = r["low"], r["high"]
    la0, ln0, la1, ln1 = lo["latitude"], lo["longitude"], hi["latitude"], hi["longitude"]
    mla, mln = (la0 + la1) / 2, (ln0 + ln1) / 2
    ret = lambda a, b, c, d: {"low": {"latitude": a, "longitude": b}, "high": {"latitude": c, "longitude": d}}
    return [ret(la0, ln0, mla, mln), ret(la0, mln, mla, ln1), ret(mla, ln0, la1, mln), ret(mla, mln, la1, ln1)]


def chave_celula(termo, r):
    s = f"{termo.lower().strip()}|" + ",".join(
        f"{v:.5f}" for v in (r["low"]["latitude"], r["low"]["longitude"], r["high"]["latitude"], r["high"]["longitude"]))
    return "grade|" + hashlib.sha1(s.encode()).hexdigest()


def buscar_grade(cliente, termo, area, profundidade=3, max_chamadas=100,
                 cache_get=None, cache_set=None, progresso=None, cancelado=None):
    """
    Varre a área em quadrantes. Quadrante que volta com 60 resultados (cheio)
    é dividido em 4 e buscado de novo, até a profundidade máxima.
    Retorna (dict place_id -> lugar, estourou_limite).
    """
    fila = [(area, 0)]
    achados, celulas = {}, 0
    while fila:
        if cancelado and cancelado():
            break
        if cliente.chamadas >= max_chamadas:
            return achados, True
        ret, prof = fila.pop(0)
        chave = chave_celula(termo, ret)
        lugares = cache_get(chave) if cache_get else None
        if lugares is None:
            lugares, completo = cliente.buscar_texto(termo, ret, cancelado=cancelado)
            if completo and cache_set:
                cache_set(chave, lugares)
        for p in lugares:
            achados.setdefault(p["id"], p)
        celulas += 1
        if len(lugares) >= LIMITE_POR_BUSCA and prof < profundidade:
            fila.extend((r, prof + 1) for r in dividir(ret))
        if progresso:
            progresso(celulas=celulas, pendentes=len(fila), encontrados=len(achados))
    return achados, False


# ---------------------------------------------------------------- linhas
def montar_linha(lugar, termo):
    tel = lugar.get("nationalPhoneNumber") or lugar.get("internationalPhoneNumber") or ""
    e164, tipo = normalizar_telefone_br(lugar.get("internationalPhoneNumber") or tel)
    loc = lugar.get("location") or {}
    celular = tipo == "Celular"
    return {
        "place_id": lugar.get("id", ""),
        "nome": (lugar.get("displayName") or {}).get("text", ""),
        "telefone": tel, "e164": e164, "tipo": tipo,
        "whatsapp": e164 if celular else "",
        "whatsapp_origem": ORIGEM_CELULAR if celular else "",
        "endereco": lugar.get("formattedAddress", ""),
        "site": lugar.get("websiteUri", ""),
        "email": "", "instagram": "",
        "avaliacao": lugar.get("rating"), "n_avaliacoes": lugar.get("userRatingCount"),
        "maps": lugar.get("googleMapsUri", ""),
        "lat": loc.get("latitude"), "lng": loc.get("longitude"),
        "business_status": lugar.get("businessStatus", "OPERATIONAL"),
        "termo": termo, "status": "Novo", "notas": "",
    }


def _valor(linha, chave):
    if chave == "link_whatsapp":
        return "" if linha.get("bloqueado") else link_whatsapp(linha.get("whatsapp", ""))
    v = linha.get(chave, "")
    return "" if v is None else v


def salvar_csv(linhas, destino):
    """destino: caminho ou objeto de texto."""
    abrir = isinstance(destino, str)
    f = open(destino, "w", newline="", encoding="utf-8-sig") if abrir else destino
    w = csv.writer(f, delimiter=";")
    w.writerow([c for c, _ in COLUNAS])
    for l in linhas:
        w.writerow([_valor(l, k) for _, k in COLUNAS])
    if abrir:
        f.close()


def salvar_xlsx(linhas, destino):
    """destino: caminho ou BytesIO. Retorna False se openpyxl não estiver instalado."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Contatos"
    ws.append([c for c, _ in COLUNAS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F6F43")
    for l in linhas:
        ws.append([_valor(l, k) for _, k in COLUNAS])
    idx = {k: i for i, (_, k) in enumerate(COLUNAS)}
    forte, leve = PatternFill("solid", fgColor="C6EBD3"), PatternFill("solid", fgColor="EEF8F1")
    for row in ws.iter_rows(min_row=2):
        origem = row[idx["whatsapp_origem"]].value
        if origem:
            for c in row:
                c.fill = forte if origem == ORIGEM_SITE else leve
        for k in ("link_whatsapp", "site", "maps"):
            c = row[idx[k]]
            if c.value:
                c.hyperlink = c.value
                c.font = Font(color="0563C1", underline="single")
    larguras = {"nome": 34, "whatsapp": 17, "whatsapp_origem": 18, "link_whatsapp": 28, "telefone": 17,
                "tipo": 12, "email": 28, "instagram": 18, "endereco": 48, "site": 32, "maps": 28, "notas": 30}
    for i, (_, k) in enumerate(COLUNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(k, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(destino)
    return True


# ---------------------------------------------------------------- modo demonstração
class DemoPlaces:
    """Simula a Places API com empresas fictícias espalhadas pela área, para testar sem chave."""
    AREA = {"low": {"latitude": -23.50, "longitude": -46.60}, "high": {"latitude": -23.38, "longitude": -46.40}}
    NOMES = ["Sorriso", "Vida", "Central", "Bem Estar", "Premium", "Família", "Saúde", "Prime", "Nova Era",
             "Total", "Real", "Ideal", "Mais", "Bela Vista", "Jardim", "Popular"]
    RUAS = ["Av. Tiradentes", "Rua Dom Pedro II", "Av. Paulo Faccini", "Rua Sete de Setembro",
            "Av. Monteiro Lobato", "Av. Salgado Filho", "Rua Capitão Gabriel"]

    def __init__(self):
        self._cache = {}

    def _empresas(self, termo):
        termo = termo.split(" em ")[0].strip().lower()
        if termo not in self._cache:
            rng = random.Random(termo)
            lo, hi = self.AREA["low"], self.AREA["high"]
            lista = []
            for i in range(rng.randint(160, 260)):
                sorte = rng.random()
                tel = (f"(11) 9{rng.randint(1000, 9999)}-{rng.randint(1000, 9999)}" if sorte < 0.55 else
                       f"(11) 2{rng.randint(100, 999)}-{rng.randint(1000, 9999)}" if sorte < 0.9 else "")
                e164, _ = normalizar_telefone_br(tel)
                lista.append({
                    "id": f"demo-{hashlib.md5(termo.encode()).hexdigest()[:6]}-{i}",
                    "displayName": {"text": f"{termo.title()} {rng.choice(self.NOMES)} {i + 1}"},
                    "nationalPhoneNumber": tel, "internationalPhoneNumber": e164,
                    "formattedAddress": f"{rng.choice(self.RUAS)}, {rng.randint(10, 3000)} - Guarulhos, SP",
                    "websiteUri": f"https://demo-{i}.exemplo.com.br" if rng.random() < 0.55 else "",
                    "googleMapsUri": "https://maps.google.com",
                    "rating": round(rng.uniform(3.4, 5.0), 1), "userRatingCount": rng.randint(2, 900),
                    "businessStatus": "OPERATIONAL" if rng.random() > 0.04 else "CLOSED_PERMANENTLY",
                    "location": {"latitude": rng.uniform(lo["latitude"], hi["latitude"]),
                                 "longitude": rng.uniform(lo["longitude"], hi["longitude"])},
                })
            lista.sort(key=lambda p: -(p["rating"] * p["userRatingCount"]))
            self._cache[termo] = lista
        return self._cache[termo]

    def search(self, corpo, mascara):
        if "viewport" in mascara:
            return {"places": [{"id": "demo-cidade", "displayName": {"text": corpo["textQuery"]},
                                "viewport": self.AREA}]}
        lista = self._empresas(corpo["textQuery"])
        r = (corpo.get("locationRestriction") or {}).get("rectangle")
        if r:
            lista = [p for p in lista
                     if r["low"]["latitude"] <= p["location"]["latitude"] <= r["high"]["latitude"]
                     and r["low"]["longitude"] <= p["location"]["longitude"] <= r["high"]["longitude"]]
        ini = int(corpo.get("pageToken") or 0)
        fim = min(ini + 20, len(lista), LIMITE_POR_BUSCA)
        d = {"places": lista[ini:fim]}
        if fim < min(len(lista), LIMITE_POR_BUSCA):
            d["nextPageToken"] = str(fim)
        return d


# ---------------------------------------------------------------- linha de comando
def main():
    import os
    p = argparse.ArgumentParser(description="Extrai contatos de empresas do Google Maps (Places API).")
    p.add_argument("termos", help='Nicho(s), vários separados por ";"')
    p.add_argument("cidade", help='Cidade/região, ex.: "Guarulhos SP"')
    p.add_argument("--grade", action="store_true", help="Cobertura total: divide a cidade em quadrantes")
    p.add_argument("--profundidade", type=int, default=3, help="Níveis de subdivisão da grade (1-4)")
    p.add_argument("--max-chamadas", type=int, default=100, help="Limite de chamadas à API")
    p.add_argument("--sites", action="store_true", help="Procurar WhatsApp/e-mail no site das empresas")
    p.add_argument("--saida", default="contatos")
    p.add_argument("--incluir-fechados", action="store_true")
    a = p.parse_args()

    demo = os.environ.get("DEMO") == "1"
    chave = os.environ.get("GOOGLE_PLACES_API_KEY", "")
    if not chave and not demo:
        sys.exit("Defina GOOGLE_PLACES_API_KEY (ou DEMO=1 para testar).")
    cli = PlacesClient(chave, demo=demo)
    coletados = {}
    try:
        area = cli.area_da_cidade(a.cidade) if a.grade else None
        for termo in [t.strip() for t in a.termos.split(";") if t.strip()]:
            print(f"Buscando: {termo} ...")
            if a.grade:
                achados, estourou = buscar_grade(cli, termo, area, a.profundidade, a.max_chamadas)
                if estourou:
                    print("  Limite de chamadas atingido; resultado parcial.")
            else:
                lista, _ = cli.buscar_texto(f"{termo} em {a.cidade}")
                achados = {x["id"]: x for x in lista}
            for pid, lugar in achados.items():
                coletados.setdefault(pid, (lugar, termo))
    except ErroAPI as e:
        sys.exit(str(e))

    linhas = [montar_linha(l, t) for l, t in coletados.values()
              if a.incluir_fechados or l.get("businessStatus", "OPERATIONAL") == "OPERATIONAL"]
    if a.sites:
        from concurrent.futures import ThreadPoolExecutor
        import enriquecer
        com_site = [l for l in linhas if l["site"]]
        print(f"Procurando WhatsApp em {len(com_site)} sites ...")
        with ThreadPoolExecutor(8) as ex:
            for l, d in zip(com_site, ex.map(lambda x: enriquecer.enriquecer_site(x["site"], demo=demo), com_site)):
                enriquecer.aplicar(l, d)
    linhas.sort(key=lambda l: (not l["whatsapp"], l["nome"].lower()))
    salvar_csv(linhas, f"{a.saida}.csv")
    ok = salvar_xlsx(linhas, f"{a.saida}.xlsx")
    tw = sum(bool(l["whatsapp"]) for l in linhas)
    print(f"\n{len(linhas)} empresas, {tw} com WhatsApp. {cli.chamadas} chamadas à API.")
    print(f"Arquivos: {a.saida}.csv" + (f" e {a.saida}.xlsx" if ok else ""))


if __name__ == "__main__":
    main()
